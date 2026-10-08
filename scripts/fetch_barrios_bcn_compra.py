"""Fetch Barcelona city/district/barrio registered home sales (Registradores).

The Generalitat publishes quarterly workbooks from Colegio de Registradores
records: transactions (new-free, new-protected, used, total), mean floor area,
mean total price (thousands of €) and mean price per built m², for Barcelona +
10 districts + 73 barris. Year pages carry BCN_trimestral_YYYY.xlsx for 2018,
2019 and 2021–2024 (2017 predates the series; 2020 returns 404); 2025 and 2026
use the name Trimestrals_Barcelona_YYYY.xlsx with extra accumulated sheets that
this fetcher skips. Only pure quarter sheets (NtYY) are parsed.

Prices/areas are unpublished where fewer than three contracts were registered
and appear as zero: stored as null, never zero. Transaction counts keep real
zeros. The city total includes records that could not be geolocated, so it is
not the sum of districts/barris (stated in the workbook footnote).
"""

from __future__ import annotations

import re
import sys
import urllib.request
from datetime import date
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from spanish_housing import manifest  # noqa: E402
from spanish_housing.data_paths import RAW  # noqa: E402

BASE = (
    "https://habitatge.gencat.cat/web/.content/home/dades/estadistiques/"
    "01_Estadistiques_de_construccio_i_mercat_immobiliari/02_Compravenda_i_preu_de_venda/"
    "02_Compravendes_d_habitatges_registrades_i_el_preu_de_venda"
)
FILES = {
    2018: "2018/BCN_trimestral_2018.xlsx",
    2019: "2019/BCN_trimestral_2019.xlsx",
    2021: "2021/BCN_trimestral_2021.xlsx",
    2022: "2022/BCN_trimestral_2022.xlsx",
    2023: "2023/BCN_trimestral_2023.xlsx",
    2024: "2024/BCN_trimestral_2024.xlsx",
    2025: "2025/Trimestrals_Barcelona_2025.xlsx",
    2026: "2026/Trimestrals_Barcelona_2026.xlsx",
}
FULL_YEARS = [2018, 2019, 2021, 2022, 2023, 2024, 2025]
RAW_PARQUET = RAW / "parquet" / "barrios_bcn_compraventes.parquet"
PUBLISHER = (
    "Generalitat de Catalunya, Secretaria d'Habitatge (Colegio de Registradores de la Propiedad)"
)
EXPECTED_HEADER = [
    "Habitatges nous lliures",
    "Hab. nous protegits",
    "Habitatge usat",
    "Total",
    "Habitatges nous lliures",
    "Hab. nous protegits",
    "Habitatge usat",
    "Total",
    "Habitatge nou",
    "Habitatge usat",
    "Total",
    "Habitatge nou",
    "Habitatge usat",
    "Total",
]
NO_DATA = {"nd", "n.d.", "n.d", "-", "...", ""}
# Publisher renamed B11 in 2025 ('el Poble Sec - Parc Montjuïc' → 'el Poble Sec',
# same geography). Canonicalize to the current label; fail on anything new.
NOM_ALIAS = {
    "el Poble Sec - Parc Montjuïc": "el Poble Sec",
    "la Marina del Prat Vermell - Zona Franca": "la Marina del Prat Vermell",
}


def download(year: int, filename: str) -> Path:
    url = f"{BASE}/{filename}"
    raw_path = RAW / f"bcn_compra_{Path(filename).name}"
    req = urllib.request.Request(url, headers={"User-Agent": "housing-data-analysis/1.0"})
    try:
        with urllib.request.urlopen(req, timeout=300) as response:  # noqa: S310 (pinned gencat host)
            contents = response.read()
    except Exception as exc:
        raise SystemExit(f"compravendes {year}: download failed: {exc}") from exc
    if not contents.startswith(b"PK"):
        raise SystemExit(f"compravendes {year}: unexpected payload (not a ZIP/XLSX)")
    tmp = raw_path.with_suffix(raw_path.suffix + ".tmp")
    tmp.write_bytes(contents)
    tmp.replace(raw_path)
    manifest.record(
        str(raw_path.relative_to(Path.cwd())),
        {
            "url": url,
            "publisher": PUBLISHER,
            "accessed": date.today().isoformat(),
            "note": f"Registradores Barcelona sales {year}; quarter sheets pinned by fetcher",
        },
    )
    return raw_path


def coerce_count(value: object) -> int | None:
    if value is None or (
        isinstance(value, str)
        and value.strip().lower().rstrip(".") in {v.rstrip(".") for v in NO_DATA}
    ):
        return None
    return int(value)


def coerce_mean(value: object, what: str) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        if value.strip().lower().rstrip(".") in {v.rstrip(".") for v in NO_DATA}:
            return None
        raise SystemExit(f"unexpected string in {what}: {value!r}")
    if value == 0:
        return None  # suppressed (<3 contracts) or no transactions: never a real zero
    return float(value)


def parse_workbook(year: int, path: Path) -> list[dict]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    quarter_sheets = [s for s in book.sheetnames if re.fullmatch(r"[1-4]t\d{2}", s)]
    if not quarter_sheets:
        raise SystemExit(f"{path.name}: no quarter sheets in {book.sheetnames}")
    rows: list[dict] = []
    seen: set[int] = set()
    for sheet in quarter_sheets:
        match = re.fullmatch(r"([1-4])t(\d{2})", sheet)
        assert match is not None
        quarter = int(match[1])
        sheet_year = 2000 + int(match[2])
        if sheet_year != year:
            raise SystemExit(f"{path.name}: sheet {sheet} does not match year {year}")
        seen.add(quarter)
        data = list(book[sheet].iter_rows(values_only=True))
        header_idx = next(i for i, r in enumerate(data) if r[0] == "Codi")
        header = [str(c).strip() for c in data[header_idx][2:] if c is not None]
        if header[: len(EXPECTED_HEADER)] != EXPECTED_HEADER:
            raise SystemExit(f"{path.name}/{sheet}: header changed: {header}")
        extra = header[len(EXPECTED_HEADER) :]
        if extra and extra != [
            "Habitatge nou",
            "Habitatge usat",
            "Habitatge nou",
            "Habitatge usat",
        ]:
            raise SystemExit(f"{path.name}/{sheet}: extra columns changed: {extra}")
        ambit = ""
        for row in data[header_idx + 1 :]:
            label = str(row[1]).strip().rstrip("*") if row[1] is not None else ""
            if label == "Barcelona":
                ambit, code, name = "ciutat", "BCN", "Barcelona"
            elif label in ("Districtes municipals", "Barris", "Barris (1)"):
                ambit = "districte" if label.startswith("Districtes") else "barri"
                continue
            elif row[0] is None or not label or label.startswith("Font:"):
                continue
            elif ambit in ("districte", "barri"):
                prefix = "D" if ambit == "districte" else "B"
                code = f"{prefix}{int(row[0]):02d}"
                name = re.sub(r"^\d+\.\s*", "", label).strip()
            else:
                if label.startswith("* Les dades"):
                    continue
                raise SystemExit(f"{path.name}/{sheet}: row outside section: {row[0:2]}")
            trx = [coerce_count(row[c]) for c in range(2, 6)]
            sup = [coerce_mean(row[c], "superficie") for c in range(7, 11)]
            price = [coerce_mean(row[c], "preu total") for c in range(12, 15)]
            eur_m2 = [coerce_mean(row[c], "eur/m2") for c in range(16, 19)]
            max_min: list[float | None] = []
            if extra:
                max_min = [coerce_mean(row[c], "eur/m2 max/min") for c in range(20, 24)]
            rows.append(
                {
                    "ambit": ambit,
                    "codi": code,
                    "nom": name,
                    "anyo": year,
                    "trimestre": quarter,
                    "trx_nou_lliure": trx[0],
                    "trx_nou_protegit": trx[1],
                    "trx_usat": trx[2],
                    "trx_total": trx[3],
                    "sup_nou_lliure": sup[0],
                    "sup_nou_protegit": sup[1],
                    "sup_usat": sup[2],
                    "sup_total": sup[3],
                    "preu_nou": price[0],
                    "preu_usat": price[1],
                    "preu_total": price[2],
                    "eur_m2_nou": eur_m2[0],
                    "eur_m2_usat": eur_m2[1],
                    "eur_m2_total": eur_m2[2],
                    "eur_m2_nou_max": max_min[0] if max_min else None,
                    "eur_m2_usat_max": max_min[1] if max_min else None,
                    "eur_m2_nou_min": max_min[2] if max_min else None,
                    "eur_m2_usat_min": max_min[3] if max_min else None,
                }
            )
    if year in FULL_YEARS and seen != {1, 2, 3, 4}:
        raise SystemExit(f"{path.name}: incomplete quarters {sorted(seen)}")
    if year == 2026 and not seen:
        raise SystemExit(f"{path.name}: no quarters published yet")
    return rows


def main() -> None:
    all_rows: list[dict] = []
    for year, filename in FILES.items():
        all_rows.extend(parse_workbook(year, download(year, filename)))
    geos = {(r["ambit"], r["codi"]) for r in all_rows}
    barris = sorted(c for (a, c) in geos if a == "barri")
    if barris != [f"B{i:02d}" for i in range(1, 74)]:
        raise SystemExit(f"barri code set changed: {barris}")
    if len(geos) != 84:
        raise SystemExit(f"area count changed: {len(geos)}")
    names = {}
    for r in all_rows:
        key = (r["ambit"], r["codi"])
        if r["nom"] in NOM_ALIAS:
            r["nom"] = NOM_ALIAS[r["nom"]]
        if key in names and names[key] != r["nom"]:
            raise SystemExit(f"label mismatch for {key}: {names[key]!r} vs {r['nom']!r}")
        names[key] = r["nom"]
    parquet = Path(str(RAW_PARQUET.relative_to(Path.cwd())))
    parquet.parent.mkdir(parents=True, exist_ok=True)
    tmp = parquet.with_suffix(parquet.suffix + ".tmp")
    pq.write_table(pa.Table.from_pylist(all_rows), tmp)
    tmp.replace(parquet)
    manifest.record(
        str(RAW_PARQUET.relative_to(Path.cwd())),
        {
            "url": "; ".join(f"{BASE}/{f}" for f in FILES.values()),
            "publisher": PUBLISHER,
            "accessed": date.today().isoformat(),
            "note": "Registradores quarterly sales; zero prices/areas stored as null "
            "(suppressed <3 contracts); city total includes non-geolocated records",
        },
    )
    full = [r for r in all_rows if r["anyo"] in FULL_YEARS]
    print(f"Barcelona compravendes: {len(all_rows)} rows")
    print(
        "nulls:",
        {
            "eur_m2_total": sum(r["eur_m2_total"] is None for r in all_rows),
            "full_year_rows": len(full),
        },
    )


if __name__ == "__main__":
    main()
