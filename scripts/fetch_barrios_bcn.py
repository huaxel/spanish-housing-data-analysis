"""Fetch Barcelona city/district/barrio rental statistics (INCASÒL deposits).

The Generalitat's housing statistics service publishes 8 XLSX workbooks with
the statistical exploitation of INCASÒL rental deposits (contracts filed):
contract counts, mean contractual rent (€/month), mean rent per m² (€/m²/month)
and mean floor area (m²). Annual files carry Barcelona + 10 districts for
2000–2025 and 73 barris for 2013–2025. Quarterly files carry one sheet per
year: districts only for 2000–2013, plus 73 barris from 2014 on; the latest
year sheet holds only the published quarters so far.

Cells for areas with fewer than six registered contracts are unpublished
(null, never zero). These are filed-contract records, not asking prices.
"""

from __future__ import annotations

import re
import sys
import urllib.request
from datetime import date
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from spanish_housing import manifest  # noqa: E402
from spanish_housing.data_paths import RAW  # noqa: E402

BASE = (
    "https://habitatge.gencat.cat/web/.content/home/dades/estadistiques/"
    "01_Estadistiques_de_construccio_i_mercat_immobiliari/03_Mercat_de_lloguer/"
    "03_Lloguers_Barcelona_per_districtes_i_barris"
)
FILES = {
    "contractes_anual": "anual_bcn_contractes.xlsx",
    "lloguer_anual": "anual_bcn_lloguer.xlsx",
    "lloguer_m2_anual": "anual_bcn_lloguer_m2.xlsx",
    "superficie_anual": "anual_bcn_sup.xlsx",
    "contractes_trimestral": "trimestral_bcn_contractes.xlsx",
    "lloguer_trimestral": "trimestral_bcn_lloguer.xlsx",
    "lloguer_m2_trimestral": "trimestral_bcn_lloguer_m2.xlsx",
    "superficie_trimestral": "trimestral_bcn_sup.xlsx",
}
MEASURES = ("contractes", "lloguer", "lloguer_m2", "superficie")
RAW_PARQUET_ANUAL = RAW / "parquet" / "barrios_bcn_lloguer_anual.parquet"
RAW_PARQUET_TRIMESTRAL = RAW / "parquet" / "barrios_bcn_lloguer_trimestral.parquet"
PUBLISHER = (
    "Generalitat de Catalunya, Servei d'Estudis i Documentació d'Habitatge / "
    "Secretaria d'Habitatge (INCASÒL dipòsits de fiances)"
)
SOURCE_LINES = {
    "Font: Servei d'Estudis i Documentació d'Habitatge, a partir de les fiances "
    "de lloguer dipositades a l'INCASÒL.",
    "Font: Secretaria d'Habitatge i Millora Urbana, a partir de les fiances "
    "de lloguer dipositades a l'INCASÒL.",
}


def download(name: str, filename: str) -> Path:
    url = f"{BASE}/{filename}"
    raw_path = RAW / f"bcn_lloguers_{filename}"
    req = urllib.request.Request(url, headers={"User-Agent": "housing-data-analysis/1.0"})
    with urllib.request.urlopen(req, timeout=300) as response:  # noqa: S310 (pinned gencat host)
        contents = response.read()
    if not contents.startswith(b"PK"):
        raise SystemExit(f"{name}: unexpected payload (not a ZIP/XLSX)")
    tmp = raw_path.with_suffix(raw_path.suffix + ".tmp")
    tmp.write_bytes(contents)
    tmp.replace(raw_path)
    manifest.record(
        str(raw_path.relative_to(Path.cwd())),
        {
            "url": url,
            "publisher": PUBLISHER,
            "accessed": date.today().isoformat(),
            "note": "Generalitat Barcelona rents workbook; sheets/years pinned by fetcher",
        },
    )
    return raw_path


def clean_district_name(name: str) -> str:
    return re.sub(r"^\d+\.\s*", "", name).strip()


# The rent workbooks keep the pre-2025 long labels for B11/B12 while the
# sales workbooks use the current short ones (same geography). Canonicalize
# to the short labels; fail on anything new.
NOM_ALIAS = {
    "el Poble Sec - AEI Parc Montjuïc": "el Poble Sec",
    "la Marina del Prat Vermell - AEI Zona Franca": "la Marina del Prat Vermell",
}


NO_DATA = {"nd", "n.d.", "-", "...", ""}


def coerce(value: object, measure: str) -> float | int | None:
    if value is None or (
        isinstance(value, str)
        and value.strip().lower().rstrip(".") in {v.rstrip(".") for v in NO_DATA}
    ):
        return None
    return int(value) if measure == "contractes" else float(value)


def section_row(label: str) -> str | None:
    if label.startswith("Districtes"):
        return "districte"
    if label == "Barris (1)":
        return "barri"
    return None


def parse_annual(
    path: Path, measure: str
) -> tuple[dict[tuple[str, str, int], float | int | None], dict[tuple[str, str], str]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(book.sheetnames) != 1:
        raise SystemExit(f"{path.name}: unexpected sheets {book.sheetnames}")
    rows = list(book[book.sheetnames[0]].iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r[0] == "Codi")
    header = rows[header_idx]
    years = [int(str(cell).strip()) for cell in header[2:] if cell is not None]
    if sorted(years) != list(range(2000, 2026)):
        raise SystemExit(f"{path.name}: unexpected year columns {years}")
    values: dict[tuple[str, str, int], float | int | None] = {}
    names: dict[tuple[str, str], str] = {}
    ambit = ""
    for row in rows[header_idx + 1 :]:
        label = str(row[1]).strip() if row[1] is not None else ""
        if label == "Barcelona":
            ambit, code, name = "ciutat", "BCN", "Barcelona"
        elif section_row(label) is not None:
            ambit = section_row(label) or ""
            continue
        elif row[0] is None or label in SOURCE_LINES or not label:
            continue
        elif ambit == "districte":
            code, name = f"D{int(row[0]):02d}", clean_district_name(label)
        elif ambit == "barri":
            code, name = f"B{int(row[0]):02d}", label.strip()
        else:
            raise SystemExit(f"{path.name}: data row outside a section: {row[0:2]}")
        cells = list(row[2 : 2 + len(years)])
        if len(cells) != len(years):
            raise SystemExit(f"{path.name}: short row for {code} {name}")
        names[(ambit, code)] = name
        for year, cell in zip(years, cells, strict=True):
            values[(ambit, code, year)] = coerce(cell, measure)
    return values, names


def parse_quarterly(
    path: Path, measure: str
) -> tuple[dict[tuple[str, str, int, int], float | int | None], dict[tuple[str, str], str]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    expected = [str(y) for y in range(2026, 1999, -1)]
    if book.sheetnames != expected:
        raise SystemExit(f"{path.name}: unexpected sheets {book.sheetnames}")
    values: dict[tuple[str, str, int, int], float | int | None] = {}
    names: dict[tuple[str, str], str] = {}
    for sheet in book.sheetnames:
        year = int(sheet)
        rows = list(book[sheet].iter_rows(values_only=True))
        new_style = any(r[0] == "Codi" for r in rows)
        if new_style:
            header_idx = next(i for i, r in enumerate(rows) if r[0] == "Codi")
            quarters = [str(rows[header_idx][c]).strip() for c in range(2, 6)]
            if quarters != ["I", "II", "III", "IV"]:
                raise SystemExit(f"{path.name}/{sheet}: unexpected quarters {quarters}")
            col0, ncols = 2, 4
        else:
            header_idx = next(i for i, r in enumerate(rows) if r[1] == "I" and r[2] == "II")
            col0, ncols = 1, 4
        ambit = ""
        for row in rows[header_idx + 1 :]:
            if new_style:
                label = str(row[1]).strip() if row[1] is not None else ""
                if label == "Barcelona":
                    ambit, code, name = "ciutat", "BCN", "Barcelona"
                elif section_row(label) is not None:
                    ambit = section_row(label) or ""
                    continue
                elif row[0] is None or label in SOURCE_LINES or not label:
                    continue
                elif ambit == "districte":
                    code, name = f"D{int(row[0]):02d}", clean_district_name(label)
                elif ambit == "barri":
                    code, name = f"B{int(row[0]):02d}", label.strip()
                else:
                    raise SystemExit(f"{path.name}/{sheet}: row outside section: {row[0:2]}")
            else:
                label = str(row[0]).strip() if row[0] is not None else ""
                if not label or label in SOURCE_LINES or label.startswith("Font:"):
                    continue
                if label == "Barcelona":
                    ambit, code, name = "ciutat", "BCN", "Barcelona"
                    cells = list(row[col0 : col0 + ncols])
                    if len(cells) != ncols:
                        raise SystemExit(f"{path.name}/{sheet}: short row for {code} {name}")
                    names[(ambit, code)] = name
                    for quarter, cell in enumerate(cells, start=1):
                        values[(ambit, code, year, quarter)] = coerce(cell, measure)
                    continue
                match = re.fullmatch(r"(\d+)\.\s*(.*)", label)
                if not match:
                    raise SystemExit(f"{path.name}/{sheet}: unexpected row {label!r}")
                ambit, code, name = "districte", f"D{int(match[1]):02d}", match[2].strip()
            cells = list(row[col0 : col0 + ncols])
            if len(cells) != ncols:
                raise SystemExit(f"{path.name}/{sheet}: short row for {code} {name}")
            names[(ambit, code)] = name
            for quarter, cell in enumerate(cells, start=1):
                values[(ambit, code, year, quarter)] = coerce(cell, measure)
    return values, names


def join_measures(
    parsed: dict[str, tuple[dict, dict]],
) -> tuple[list[dict], dict[tuple[str, str], str]]:
    keys: set[tuple] = set()
    for vals, _ in parsed.values():
        keys |= set(vals)
    names: dict[tuple[str, str], str] = {}
    for _, (_, measure_names) in parsed.items():
        for geo, name in measure_names.items():
            name = NOM_ALIAS.get(name, name)
            if geo in names and names[geo] != name:
                raise SystemExit(f"geography label mismatch for {geo}: {names[geo]!r} vs {name!r}")
            names[geo] = name
    rows = []
    for key in sorted(keys):
        ambit, code = key[0], key[1]
        row: dict = {"ambit": ambit, "codi": code, "nom": names[(ambit, code)]}
        if len(key) == 3:
            row["anyo"] = key[2]
            for measure in MEASURES:
                row[measure] = parsed[measure][0].get(key)
        else:
            row["anyo"], row["trimestre"] = key[2], key[3]
            for measure in MEASURES:
                row[measure] = parsed[measure][0].get(key)
        rows.append(row)
    return rows, names


def write_parquet(rows: list[dict], rel: str, note: str, urls: list[str]) -> None:
    parquet = Path(rel)
    parquet.parent.mkdir(parents=True, exist_ok=True)
    tmp = parquet.with_suffix(parquet.suffix + ".tmp")
    pq.write_table(pa.Table.from_pylist(rows), tmp)
    tmp.replace(parquet)
    manifest.record(
        rel,
        {
            "url": "; ".join(urls),
            "publisher": PUBLISHER,
            "accessed": date.today().isoformat(),
            "note": note,
        },
    )


def main() -> None:
    paths = {key: download(key, filename) for key, filename in FILES.items()}
    anual = {m: parse_annual(paths[f"{m}_anual"], m) for m in MEASURES}
    trimestral = {m: parse_quarterly(paths[f"{m}_trimestral"], m) for m in MEASURES}

    anual_rows, anual_names = join_measures(anual)
    city = [r for r in anual_rows if r["ambit"] == "ciutat"]
    districts = [r for r in anual_rows if r["ambit"] == "districte"]
    barris = [r for r in anual_rows if r["ambit"] == "barri"]
    if not (len(city) == 26 and len(districts) == 260 and len(barris) == 73 * 26):
        raise SystemExit(
            f"annual coverage changed: city={len(city)} districts={len(districts)} "
            f"barris={len(barris)}"
        )
    if any(r["contractes"] is not None for r in barris if r["anyo"] < 2013):
        raise SystemExit("annual barri history unexpectedly extends before 2013")
    if sorted({c for (_, c) in [k for k in anual_names if k[0] == "barri"]}) != [
        f"B{i:02d}" for i in range(1, 74)
    ]:
        raise SystemExit("annual barri code set changed")
    barri_years = sorted({r["anyo"] for r in barris if r["contractes"] is not None})
    if barri_years != list(range(2013, 2026)):
        raise SystemExit(f"annual barri years changed: {barri_years}")

    trim_rows, trim_names = join_measures(trimestral)
    trim_barris = [r for r in trim_rows if r["ambit"] == "barri"]
    trim_barri_years = sorted({r["anyo"] for r in trim_barris})
    if trim_barri_years != list(range(2014, 2027)):
        raise SystemExit(f"quarterly barri years changed: {trim_barri_years}")
    expected_trim = 11 * 27 * 4 + 73 * 13 * 4
    if len(trim_rows) != expected_trim:
        raise SystemExit(f"quarterly coverage changed: {len(trim_rows)} rows")

    urls = [f"{BASE}/{f}" for f in FILES.values()]
    write_parquet(
        anual_rows,
        str(RAW_PARQUET_ANUAL.relative_to(Path.cwd())),
        "INCASÒL deposits: city+districts 2000–2025, 73 barris 2013–2025; "
        "<6-contract cells null; filed contracts, not asking prices",
        urls,
    )
    write_parquet(
        trim_rows,
        str(RAW_PARQUET_TRIMESTRAL.relative_to(Path.cwd())),
        "INCASÒL deposits quarterly: districts 2000–2013, +73 barris 2014–; "
        "latest year holds published quarters only; <6-contract cells null",
        urls,
    )
    print(f"Barcelona lloguers anual: {len(anual_rows)} rows")
    print(f"Barcelona lloguers trimestral: {len(trim_rows)} rows")
    print(
        "nulls:",
        {
            "anual_contractes": sum(r["contractes"] is None for r in anual_rows),
            "anual_lloguer_m2": sum(r["lloguer_m2"] is None for r in anual_rows),
            "trim_lloguer_m2": sum(r["lloguer_m2"] is None for r in trim_rows),
        },
    )


if __name__ == "__main__":
    main()
